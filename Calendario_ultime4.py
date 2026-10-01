#!/usr/bin/env python3
"""Genera quattro giornate con 24 squadre e salva il calendario in Excel.

Vincoli garantiti:
- ogni squadra gioca una partita per giornata;
- ogni squadra gioca 2 volte in casa e 2 volte in trasferta;
- nessun accoppiamento viene ripetuto nelle quattro giornate.

Dipendenza:
    pip install openpyxl

Esempio:
    python genera_calendario_4_giornate.py --output Calendario_4_giornate.xlsx
"""

from __future__ import annotations

import argparse
from collections import defaultdict
from pathlib import Path
from typing import Callable, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


SQUADRE = [
    "A.O. Pelicancer",
    "RSM Money Launderers",
    "Montecelio Redskins",
    "Pasadena Pupils",
    "Sex Pistons",
    "NMMFSIV",
    "Varese Konige",
    "Baltimora Bats",
    "Tibur Eagles",
    "Lords of Torture",
    "Champapon Newborns",
    "Ntilikinerz",
    "WestSide Hammers",
    "Ohio Raptors",
    "Buffalo Brownies",
    "Polisportiva Calitri",
    "LongNosed Milfhunter",
    "BOOKER-T",
    "ZZ Killers",
    "No Trade Knicks",
    "Seattle 206ers",
    "Guston Rockets",
    "Bologna Boricuas",
    "DREAMCHASERS",
]


def mulberry32(seed: int) -> Callable[[], float]:
    """PRNG compatibile con quello usato per generare il file originale."""
    mask = 0xFFFFFFFF
    state = seed & mask

    def imul(a: int, b: int) -> int:
        return (a * b) & mask

    def random_value() -> float:
        nonlocal state
        state = (state + 0x6D2B79F5) & mask
        value = state
        value = imul(value ^ (value >> 15), value | 1)
        value ^= (value + imul(value ^ (value >> 7), value | 61)) & mask
        value &= mask
        return ((value ^ (value >> 14)) & mask) / 2**32

    return random_value


def mescola(elementi: Iterable, random_value: Callable[[], float]) -> list:
    risultato = list(elementi)
    for indice in range(len(risultato) - 1, 0, -1):
        altro_indice = int(random_value() * (indice + 1))
        risultato[indice], risultato[altro_indice] = (
            risultato[altro_indice],
            risultato[indice],
        )
    return risultato


def crea_giornate(squadre: list[str], seed: int) -> list[list[tuple[str, str]]]:
    if len(squadre) != 24 or len(set(squadre)) != 24:
        raise ValueError("Servono esattamente 24 nomi di squadra univoci.")

    random_value = mulberry32(seed)
    squadre_mescolate = mescola(squadre, random_value)

    # Ogni schema contiene due gare in casa (1) e due in trasferta (0).
    # Gli schemi sono accoppiati con il rispettivo complemento.
    schemi_casa = [
        [1, 1, 0, 0],
        [0, 0, 1, 1],
        [1, 0, 1, 0],
        [0, 1, 0, 1],
        [1, 0, 0, 1],
        [0, 1, 1, 0],
    ]
    gruppi = [
        {
            "schema": schema,
            "squadre": squadre_mescolate[indice * 4 : indice * 4 + 4],
        }
        for indice, schema in enumerate(schemi_casa)
    ]

    giornate: list[list[tuple[str, str]]] = [[] for _ in range(4)]
    for indice_coppia in range(3):
        gruppo_a = gruppi[indice_coppia * 2]
        gruppo_b = gruppi[indice_coppia * 2 + 1]
        rotazioni = mescola([0, 1, 2, 3], random_value)

        for indice_giornata in range(4):
            rotazione = rotazioni[indice_giornata]
            for indice_squadra in range(4):
                squadra_a = gruppo_a["squadre"][indice_squadra]
                squadra_b = gruppo_b["squadre"][(indice_squadra + rotazione) % 4]
                if gruppo_a["schema"][indice_giornata] == 1:
                    giornate[indice_giornata].append((squadra_a, squadra_b))
                else:
                    giornate[indice_giornata].append((squadra_b, squadra_a))

    return [mescola(giornata, random_value) for giornata in giornate]


def verifica_giornate(
    giornate: list[list[tuple[str, str]]], squadre: list[str]
) -> None:
    conteggi = defaultdict(lambda: {"casa": 0, "trasferta": 0, "partite": 0})
    accoppiamenti: set[tuple[str, str]] = set()

    if len(giornate) != 4:
        raise ValueError("Il calendario deve contenere quattro giornate.")

    for numero, giornata in enumerate(giornate, start=1):
        if len(giornata) != 12:
            raise ValueError(f"Giornata {numero}: devono esserci 12 partite.")

        presenti: set[str] = set()
        for casa, trasferta in giornata:
            if casa in presenti or trasferta in presenti:
                raise ValueError(f"Giornata {numero}: una squadra compare più volte.")
            presenti.update((casa, trasferta))

            conteggi[casa]["casa"] += 1
            conteggi[casa]["partite"] += 1
            conteggi[trasferta]["trasferta"] += 1
            conteggi[trasferta]["partite"] += 1

            coppia = tuple(sorted((casa, trasferta)))
            if coppia in accoppiamenti:
                raise ValueError(f"Accoppiamento ripetuto: {coppia[0]} - {coppia[1]}")
            accoppiamenti.add(coppia)

        if presenti != set(squadre):
            raise ValueError(f"Giornata {numero}: l'elenco delle squadre non è completo.")

    for squadra in squadre:
        risultato = conteggi[squadra]
        if risultato != {"casa": 2, "trasferta": 2, "partite": 4}:
            raise ValueError(f"Vincolo non rispettato per {squadra}: {risultato}")


def salva_excel(giornate: list[list[tuple[str, str]]], destinazione: Path) -> None:
    workbook = Workbook()
    foglio = workbook.active
    foglio.title = "Calendario"
    foglio.sheet_view.showGridLines = True

    riquadri = [
        ("GIORNATA 1", 1, 1),
        ("GIORNATA 2", 1, 4),
        ("GIORNATA 3", 16, 1),
        ("GIORNATA 4", 16, 4),
    ]

    bordo_grigio = Side(style="thin", color="D9D9D9")
    bordo_titolo = Side(style="thin", color="BFBFBF")
    font_corpo = Font(name="Arial", size=10, color="000000")
    font_titolo = Font(name="Arial", size=11, bold=True, color="111827")
    riempimento_titolo = PatternFill("solid", fgColor="E7E6E6")

    for indice, (titolo, riga, colonna) in enumerate(riquadri):
        foglio.merge_cells(
            start_row=riga,
            start_column=colonna,
            end_row=riga,
            end_column=colonna + 1,
        )
        cella_titolo = foglio.cell(riga, colonna, titolo)
        cella_titolo.font = font_titolo
        cella_titolo.fill = riempimento_titolo
        cella_titolo.alignment = Alignment(horizontal="center", vertical="center")

        for cella in foglio[riga][colonna - 1 : colonna + 1]:
            cella.border = Border(
                left=bordo_titolo,
                right=bordo_titolo,
                top=bordo_titolo,
                bottom=bordo_titolo,
            )

        for offset, (casa, trasferta) in enumerate(giornate[indice], start=1):
            for posizione, valore in enumerate((casa, trasferta)):
                cella = foglio.cell(riga + offset, colonna + posizione, valore)
                cella.font = font_corpo
                cella.alignment = Alignment(horizontal="left", vertical="center")
                cella.border = Border(
                    left=bordo_grigio,
                    right=bordo_grigio,
                    top=bordo_grigio,
                    bottom=bordo_grigio,
                )

    for colonna, larghezza in {1: 27, 2: 27, 3: 3, 4: 27, 5: 27}.items():
        foglio.column_dimensions[get_column_letter(colonna)].width = larghezza

    for riga in (1, 16):
        foglio.row_dimensions[riga].height = 23
    for riga in list(range(2, 14)) + list(range(17, 29)):
        foglio.row_dimensions[riga].height = 20
    for riga in (14, 15):
        foglio.row_dimensions[riga].height = 8

    destinazione.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destinazione)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("Calendario_4_giornate.xlsx"),
        help="Percorso del file Excel da creare.",
    )
    parser.add_argument(
        "--seed",
        type=int,
        default=20261001,
        help="Seed del sorteggio. Cambialo per ottenere un calendario diverso.",
    )
    argomenti = parser.parse_args()

    giornate = crea_giornate(SQUADRE, argomenti.seed)
    verifica_giornate(giornate, SQUADRE)
    salva_excel(giornate, argomenti.output)
    print(f"File creato: {argomenti.output.resolve()}")


if __name__ == "__main__":
    main()
